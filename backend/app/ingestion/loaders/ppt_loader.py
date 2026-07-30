from pathlib import Path
from pptx import Presentation
from .base_loader import BaseLoader


class PPTLoader(BaseLoader):

    def load(self, path: Path) -> str:
        prs = Presentation(str(path))
        sections = []

        for i, slide in enumerate(prs.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    parts.append(shape.text.strip())
            if parts:
                sections.append(f"## Slide {i}\n\n" + "\n".join(parts))

        return "\n\n".join(sections)
